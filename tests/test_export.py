"""Integration tests for multi-tab export and Google Sheets formatting."""

import tempfile
import os
import gc
import openpyxl
import pytest

from gseim.database.repository import GSEIMRepository
from gseim.export.excel_mirror import export_to_excel
from gseim.config.settings import SHEET_TABS

@pytest.fixture
def populated_repo():
    fd, path = tempfile.mkstemp(suffix=".sqlite3")
    os.close(fd)
    repo = GSEIMRepository(db_path=path)

    # Insert sample source
    repo.save_opportunity({
        "organization": "Test Accelerator",
        "country": "Singapore",
        "region": "Southeast Asia",
        "type": "Accelerator",
        "opportunity": "2027 ClimateTech Accelerator",
        "sector": "ClimateTech",
        "startup_stage": "Seed",
        "eligibility": "Startups incorporated in Singapore",
        "application_status": "Open",
        "deadline": "15-Nov-2026",
        "deadline_iso": "2026-11-15",
        "days_remaining": 47,
        "source_url": "https://sg.org/climate",
        "source_type": "Official Website",
        "relevance_score": 14,
        "summary": "Cohort program."
    })

    repo.save_job({
        "organization": "XYZ Innovation Hub",
        "country": "UK",
        "region": "Europe",
        "job_title": "Program Manager",
        "function": "Incubation",
        "seniority": "Mid-Senior",
        "employment": "Full-time",
        "location": "London",
        "remote": "No",
        "salary": "£50,000–£60,000",
        "closing_date": "20-Oct-2026",
        "closing_iso": "2026-10-20",
        "days_remaining": 21,
        "job_url": "https://xyz.org/careers/pm",
        "source": "Official Website",
        "relevance_score": 18,
        "summary": "Incubation manager for cohort."
    })

    repo.save_ecosystem_update({
        "date": "29-09-2026",
        "organization": "Government",
        "country": "India",
        "category": "Government Fund",
        "title": "Government launches ₹500 crore startup fund",
        "summary": "New deeptech fund established.",
        "url": "https://gov.in/fund"
    })

    yield repo
    del repo
    gc.collect()
    try:
        if os.path.exists(path):
            os.remove(path)
    except PermissionError:
        pass

def test_excel_export_all_10_tabs(populated_repo):
    fd, out_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)

    try:
        exported_path = export_to_excel(populated_repo, output_path=out_path)
        assert os.path.exists(exported_path)

        wb = openpyxl.load_workbook(exported_path)
        sheet_names = wb.sheetnames

        # Verify all 10 tabs exist
        for expected_tab in SHEET_TABS:
            assert expected_tab in sheet_names, f"Missing expected tab: {expected_tab}"

        # Verify Opportunities sheet has data
        opp_ws = wb["New Opportunities"]
        assert opp_ws.max_row >= 2  # header + at least 1 record
        header_vals = [opp_ws.cell(row=1, column=c).value for c in range(1, 19)]
        assert "Opportunity" in header_vals
        assert "Relevance Score" in header_vals

        # Verify Jobs sheet has data
        job_ws = wb["Jobs"]
        assert job_ws.max_row >= 2
        job_headers = [job_ws.cell(row=1, column=c).value for c in range(1, 17)]
        assert "Job Title" in job_headers
        assert "Function" in job_headers

        wb.close()
    finally:
        gc.collect()
        try:
            if os.path.exists(out_path):
                os.remove(out_path)
        except PermissionError:
            pass
